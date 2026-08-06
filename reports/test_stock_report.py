"""Отчёт об остатках: артикул материалов и деление на материалы и продукцию.

Отчёт долго показывал у материалов прочерк вместо артикула и цвета —
поля брались только у продукции, хотя у материала они тоже есть и
заполняются в справочнике. Ошибка сидела сразу в двух местах: на
странице и в выгрузке в Excel, потому что сборка строк была написана
дважды. Теперь она одна, и тесты держат оба выхода.
"""
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase

from warehouse.models import Material, Product, Stock, Warehouse

User = get_user_model()


class StockReportTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user(
            'kladovshchik', password='x', is_superuser=True)
        cls.warehouse = Warehouse.objects.create(
            name='Основной', type='raw')
        cls.material = Material.objects.create(
            name='Кожа', article_number='КЖ-100', color='красный',
            unit='m', category='leather', reorder_point=10)
        cls.product = Product.objects.create(
            article_number='МАК-01', name='Макивара', category='equipment',
            size='L', color='синий', cost=500, selling_price=1300,
            status='active')
        Stock.objects.create(warehouse=cls.warehouse,
                             material=cls.material, quantity=250)
        Stock.objects.create(warehouse=cls.warehouse,
                             product=cls.product, quantity=40)

    def setUp(self):
        self.client.force_login(self.user)

    def rows(self, kind=None):
        url = '/api/reports/stock/'
        if kind:
            url += f'?kind={kind}'
        return self.client.get(url).json()['строки']

    # --- Артикул и цвет материала --------------------------------------
    def test_material_article_is_shown(self):
        """Артикул материала берётся из справочника, а не прочерк."""
        row = next(r for r in self.rows() if r['тип'] == 'Материал')
        self.assertEqual(row['артикул'], 'КЖ-100')

    def test_material_color_is_shown(self):
        row = next(r for r in self.rows() if r['тип'] == 'Материал')
        self.assertEqual(row['цвет'], 'красный')

    def test_product_article_still_shown(self):
        row = next(r for r in self.rows() if r['тип'] == 'Продукция')
        self.assertEqual(row['артикул'], 'МАК-01')

    def test_material_has_no_size(self):
        """Размера у материала нет — его меряют единицей измерения."""
        row = next(r for r in self.rows() if r['тип'] == 'Материал')
        self.assertEqual(row['размер'], '—')
        self.assertEqual(row['единица'], 'м')

    def test_empty_article_becomes_dash(self):
        """Незаполненный артикул показывается прочерком, а не пустотой."""
        Material.objects.filter(pk=self.material.pk).update(article_number='')
        row = next(r for r in self.rows() if r['тип'] == 'Материал')
        self.assertEqual(row['артикул'], '—')

    # --- Деление отчёта -------------------------------------------------
    def test_kind_material_leaves_only_materials(self):
        rows = self.rows('material')
        self.assertTrue(rows)
        self.assertTrue(all(r['тип'] == 'Материал' for r in rows))

    def test_kind_product_leaves_only_products(self):
        rows = self.rows('product')
        self.assertTrue(rows)
        self.assertTrue(all(r['тип'] == 'Продукция' for r in rows))

    def test_default_shows_everything(self):
        kinds = {r['тип'] for r in self.rows()}
        self.assertEqual(kinds, {'Материал', 'Продукция'})

    def test_unknown_kind_falls_back_to_all(self):
        """Опечатка в адресе не должна выдавать пустой отчёт."""
        data = self.client.get('/api/reports/stock/?kind=ерунда').json()
        self.assertEqual(data['вид'], 'all')
        self.assertEqual(len(data['строки']), 2)

    def test_title_says_what_is_shown(self):
        for kind, title in [('material', 'Остатки материалов'),
                            ('product', 'Остатки продукции')]:
            with self.subTest(kind=kind):
                data = self.client.get(
                    f'/api/reports/stock/?kind={kind}').json()
                self.assertEqual(data['отчёт'], title)

    def test_zero_stock_not_included(self):
        Stock.objects.filter(material=self.material).update(quantity=0)
        self.assertEqual(len(self.rows('material')), 0)


class StockExportTest(TestCase):
    """Выгрузка в Excel показывает ровно то же, что и страница."""

    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user(
            'buhgalter', password='x', is_superuser=True)
        warehouse = Warehouse.objects.create(name='Основной', type='raw')
        material = Material.objects.create(
            name='Кожа', article_number='КЖ-100', color='красный',
            unit='m', category='leather', reorder_point=10)
        product = Product.objects.create(
            article_number='МАК-01', name='Макивара', category='equipment',
            size='L', color='синий', cost=500, selling_price=1300,
            status='active')
        Stock.objects.create(warehouse=warehouse, material=material,
                             quantity=250)
        Stock.objects.create(warehouse=warehouse, product=product,
                             quantity=40)

    def setUp(self):
        self.client.force_login(self.user)

    def sheet(self, kind=None):
        import openpyxl
        url = '/api/reports/stock/export/'
        if kind:
            url += f'?kind={kind}'
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        return openpyxl.load_workbook(BytesIO(response.content)).active

    def test_export_opens_as_excel(self):
        """Файл должен открываться, а не быть страницей с ошибкой.

        Выгрузка — обычное представление Django, у неё нет query_params;
        обращение к ним отдавало HTML с трассировкой под именем .xlsx.
        """
        self.assertGreater(self.sheet().max_row, 1)

    def test_export_shows_material_article(self):
        rows = list(self.sheet('material').iter_rows(min_row=2,
                                                     values_only=True))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][3], 'КЖ-100')   # столбец «Артикул»
        self.assertEqual(rows[0][5], 'красный')  # столбец «Цвет»

    def test_export_respects_kind(self):
        self.assertEqual(self.sheet('material').max_row - 1, 1)
        self.assertEqual(self.sheet('product').max_row - 1, 1)
        self.assertEqual(self.sheet().max_row - 1, 2)

    def test_sheet_named_after_selection(self):
        self.assertEqual(self.sheet('material').title, 'Материалы')
        self.assertEqual(self.sheet('product').title, 'Продукция')
        self.assertEqual(self.sheet().title, 'Остатки')

    def test_file_name_says_what_is_inside(self):
        """Три выгрузки в одной папке должны различаться по имени."""
        for kind, name in [('material', 'materials'), ('product', 'products'),
                           (None, 'stock')]:
            with self.subTest(kind=kind):
                url = '/api/reports/stock/export/'
                if kind:
                    url += f'?kind={kind}'
                disposition = self.client.get(url)['Content-Disposition']
                self.assertIn(f'filename="{name}.xlsx"', disposition)

    def test_export_requires_login(self):
        self.client.logout()
        response = self.client.get('/api/reports/stock/export/')
        self.assertIn(response.status_code, (302, 401, 403))
