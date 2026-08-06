"""
Отчёты системы складского учёта.

Формируются пять видов отчётов (см. п. 2.2.5 отчёта):
1. Остатки на дату — текущие остатки по складам и категориям
2. Движение товаров за период — приходы, расходы, корректировки
3. Расход материалов по производственным заказам
4. Результаты инвентаризаций — расхождения и недостачи
5. Перечень позиций для закупки — с остатком ниже минимума

Отчёты выводятся в JSON и экспортируются в Excel.
"""
from datetime import datetime, timedelta
from io import BytesIO
from urllib.parse import quote

from django.db.models import DecimalField, F, Sum, Value
from django.db.models.functions import Coalesce
from django.contrib.auth.decorators import login_required
from django.utils import timezone
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from warehouse.models import Material, Stock, StockMovement
from inventory.models import Inventory


#: Что показывать в отчёте об остатках. Материалы и продукция живут по
#: разным правилам: у материала нет размера, у продукции нет единицы
#: измерения кроме штук, — и смотрят на них разные люди. Снабженец
#: считает сырьё, а сбыт — готовые изделия.
STOCK_KINDS = ('all', 'material', 'product')

DASH = '—'


def read_days(value, default, low=1, high=3650):
    """Число дней из адреса страницы, приведённое к разумным границам.

    Параметр приходит из браузера, и там может оказаться что угодно:
    пусто, буквы, число в сто знаков. Раньше отчёт о движении переводил
    его в число напрямую, и «?days=abc» отдавал отказ сервера вместо
    ответа. Слишком большое значение валило перевод отдельно: такое
    число не помещается в целое, с которым работает база.
    """
    if value in (None, ''):
        return default
    try:
        number = int(value)
    except (TypeError, ValueError):
        raise ValueError(f'Неверное значение параметра days: {value!r}')
    return max(low, min(number, high))


#: Предел для номера записи. База хранит их в 64-разрядном целом, и
#: число сверх этого валит запрос ещё до отбора — с отказом сервера, а
#: не с пустым ответом.
MAX_ID = 2 ** 63 - 1


def read_id(value, name):
    """Номер записи из адреса страницы.

    Отдать нечисловой номер прямо в отбор нельзя: база отвечает
    «Field 'id' expected a number», и пользователь видит отказ сервера
    вместо ответа. Слишком большое число валится отдельно, уже в SQLite.
    Пустое значение означает «без отбора».
    """
    if value in (None, ''):
        return None
    try:
        number = int(value)
    except (TypeError, ValueError):
        raise ValueError(f'Неверное значение параметра {name}: {value!r}')
    if not (0 < number <= MAX_ID):
        raise ValueError(f'Неверное значение параметра {name}: {value!r}')
    return number


def _stock_rows(stocks):
    """Строки отчёта об остатках.

    Одна и та же сборка нужна и странице, и выгрузке в Excel. Раньше
    они были написаны порознь, и одинаковая ошибка сидела в обеих:
    артикул и цвет брались только у продукции, а у материала
    подставлялся прочерк — хотя оба поля у материала есть и заполняются
    в справочнике.
    """
    rows = []
    for stock in stocks:
        item = stock.material or stock.product
        if item is None:
            # Остаток без товара — испорченная строка, в отчёт не берём
            continue
        is_material = stock.material is not None
        rows.append({
            'склад': stock.warehouse.name,
            'тип': 'Материал' if is_material else 'Продукция',
            'наименование': item.name,
            'артикул': item.article_number or DASH,
            # Размер есть только у продукции: материал меряют не им,
            # а единицей измерения — метрами, килограммами.
            'размер': DASH if is_material else (item.size or DASH),
            'цвет': item.color or DASH,
            'количество': float(stock.quantity),
            'единица': item.get_unit_display() if is_material else 'шт.',
        })
    return rows


def _stock_queryset(request):
    """Остатки с учётом отбора по складу и виду номенклатуры.

    Читает параметры через request.GET, а не query_params: выгрузка в
    Excel — обычное представление Django, у неё query_params нет, и
    обращение к нему роняло скачивание файла.
    """
    params = request.GET
    stocks = Stock.objects.select_related('warehouse', 'material', 'product')

    warehouse_id = read_id(params.get('warehouse'), 'warehouse')
    if warehouse_id is not None:
        stocks = stocks.filter(warehouse_id=warehouse_id)

    kind = params.get('kind', 'all')
    if kind not in STOCK_KINDS:
        kind = 'all'
    if kind == 'material':
        stocks = stocks.filter(material__isnull=False)
    elif kind == 'product':
        stocks = stocks.filter(product__isnull=False)

    return stocks.filter(quantity__gt=0), kind


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def stock_report(request):
    """Остатки на дату: текущие остатки по складам.

    Параметр kind делит отчёт: material — только сырьё, product —
    только готовые изделия, all (по умолчанию) — всё вместе.
    """
    date = request.query_params.get('date', datetime.now().date())
    try:
        stocks, kind = _stock_queryset(request)
    except ValueError as e:
        return Response({'error': str(e)}, status=400)
    rows = _stock_rows(stocks)

    titles = {
        'material': 'Остатки материалов',
        'product': 'Остатки продукции',
        'all': 'Остатки на дату',
    }

    return Response({
        'отчёт': titles.get(kind, titles['all']),
        'вид': kind if kind in STOCK_KINDS else 'all',
        'дата': date,
        'всего_позиций': len(rows),
        'строки': rows,
    })


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def movement_report(request):
    """Движение товаров за период."""
    try:
        days_back = read_days(request.query_params.get('days'), 30)
    except ValueError as e:
        return Response({'error': str(e)}, status=400)
    from_date = datetime.now().date() - timedelta(days=days_back)
    to_date = datetime.now().date()
    try:
        warehouse_id = read_id(request.query_params.get('warehouse'),
                               'warehouse')
    except ValueError as e:
        return Response({'error': str(e)}, status=400)

    movements = (StockMovement.objects
                .select_related('warehouse', 'material', 'product', 'user')
                .filter(created_at__date__gte=from_date,
                        created_at__date__lte=to_date))

    if warehouse_id is not None:
        movements = movements.filter(warehouse_id=warehouse_id)

    inbound_qty = (movements.filter(movement_type='in')
                  .aggregate(Sum('quantity'))['quantity__sum'] or 0)
    outbound_qty = (movements.filter(movement_type='out')
                   .aggregate(Sum('quantity'))['quantity__sum'] or 0)

    rows = []
    for movement in movements.order_by('-created_at'):
        rows.append({
            'дата': str(movement.created_at.date()),
            'время': movement.created_at.time().isoformat(),
            'тип': movement.get_movement_type_display(),
            'товар': movement.item_name,
            'количество': float(movement.quantity),
            'склад': movement.warehouse.name,
            'пользователь': movement.user.get_full_name() if movement.user else '—',
            'документ': movement.document_id or '—',
        })

    return Response({
        'отчёт': 'Движение товаров',
        'с_даты': str(from_date),
        'по_дату': str(to_date),
        'приход_кол': float(inbound_qty),
        'расход_кол': float(outbound_qty),
        'всего_операций': len(rows),
        'строки': rows,
    })


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def reorder_report(request):
    """Перечень позиций для закупки — с остатком ниже минимума."""
    materials = (Material.objects
                 .annotate(
                     stock_qty=Coalesce(
                         Sum('stock__quantity'),
                         Value(0, output_field=DecimalField())))
                 .filter(stock_qty__lt=F('reorder_point')))

    rows = []
    for material in materials:
        shortage = material.reorder_point - material.stock_qty
        rows.append({
            'наименование': material.name,
            'категория': material.get_category_display(),
            'текущий_остаток': float(material.stock_qty),
            'минимум': float(material.reorder_point),
            'недостаток': float(shortage),
            'единица': material.get_unit_display(),
            'описание': material.description,
        })

    return Response({
        'отчёт': 'Позиции для закупки',
        'дата': str(datetime.now().date()),
        'позиций_в_дефиците': len(rows),
        'строки': rows,
    })


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def inventory_report(request):
    """Результаты инвентаризаций."""
    try:
        inventory_id = read_id(request.query_params.get('inventory'),
                               'inventory')
    except ValueError as e:
        return Response({'error': str(e)}, status=400)

    if inventory_id is not None:
        inventories = Inventory.objects.filter(pk=inventory_id)
    else:
        inventories = (Inventory.objects
                      .filter(status='completed')
                      .order_by('-completed_at')[:10])

    results = []
    for inv in inventories:
        shortages = []
        surpluses = []
        for item in inv.items.all():
            if item.difference < 0:
                shortages.append({
                    'товар': item.item_name,
                    'величина': float(abs(item.difference)),
                })
            elif item.difference > 0:
                surpluses.append({
                    'товар': item.item_name,
                    'величина': float(item.difference),
                })

        results.append({
            'номер': inv.number,
            'склад': inv.warehouse.name,
            'дата_завершения': str(inv.completed_at.date()) if inv.completed_at else None,
            'недостач': len(shortages),
            'излишков': len(surpluses),
            'недостачи': shortages[:5],
            'излишки': surpluses[:5],
        })

    return Response({
        'отчёт': 'Результаты инвентаризаций',
        'дата': str(datetime.now().date()),
        'всего_инвентаризаций': len(results),
        'результаты': results,
    })


@login_required
def stock_report_export(request):
    """Экспорт отчёта об остатках в Excel."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        from django.http import HttpResponse
        return HttpResponse('Библиотека openpyxl не установлена', status=400)

    # Данные собираются той же функцией, что и для страницы: выгрузка
    # должна показывать ровно то, что человек видел на экране.
    try:
        stocks, kind = _stock_queryset(request)
    except ValueError as e:
        from django.http import HttpResponse
        return HttpResponse(str(e), status=400, content_type='text/plain; charset=utf-8')
    rows = _stock_rows(stocks)

    sheet_titles = {
        'material': 'Материалы',
        'product': 'Продукция',
        'all': 'Остатки',
    }

    # Создать книгу
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_titles.get(kind, sheet_titles['all'])

    # Заголовок
    headers = ['Склад', 'Тип', 'Наименование', 'Артикул', 'Размер',
              'Цвет', 'Количество', 'Единица']
    worksheet.append(headers)

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='366092', end_color='366092',
                            fill_type='solid')
    for cell in worksheet[1]:
        cell.font = header_font
        cell.fill = header_fill

    # Данные
    for row in rows:
        worksheet.append([
            row['склад'], row['тип'], row['наименование'], row['артикул'],
            row['размер'], row['цвет'], row['количество'], row['единица'],
        ])

    # Установить ширину столбцов
    worksheet.column_dimensions['A'].width = 20
    worksheet.column_dimensions['B'].width = 12
    worksheet.column_dimensions['C'].width = 25
    worksheet.column_dimensions['D'].width = 12
    worksheet.column_dimensions['E'].width = 10
    worksheet.column_dimensions['F'].width = 12
    worksheet.column_dimensions['G'].width = 12
    worksheet.column_dimensions['H'].width = 12

    # Сохранить в памяти и отправить как файл (HttpResponse, не DRF Response —
    # иначе DRF пытается декодировать бинарные байты как UTF-8 и падает)
    from django.http import HttpResponse
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument'
                     '.spreadsheetml.sheet')
    # Имя файла говорит, что внутри: три выгрузки в одной папке иначе
    # неразличимы. Латиницей в filename + кириллица в filename* (RFC 5987).
    ascii_names = {'material': 'materials', 'product': 'products',
                   'all': 'stock'}
    ru_names = {'material': 'остатки-материалов',
                'product': 'остатки-продукции',
                'all': 'остатки'}
    ascii_name = ascii_names.get(kind, ascii_names['all'])
    ru_name = quote(ru_names.get(kind, ru_names['all']))
    response['Content-Disposition'] = (
        f'attachment; filename="{ascii_name}.xlsx"; '
        f"filename*=UTF-8''{ru_name}.xlsx")
    return response


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def expiry_report(request):
    """Партии с истекающим сроком годности.

    Параметр ?days= задаёт горизонт предупреждения в днях (по умолчанию 30).
    Просроченное показывается всегда, независимо от горизонта.
    """
    from warehouse.expiry import DEFAULT_WARNING_DAYS, expiring_batches

    raw_days = request.query_params.get('days')
    if raw_days in (None, ''):
        days = DEFAULT_WARNING_DAYS
    else:
        try:
            days = int(raw_days)
        except (TypeError, ValueError):
            return Response(
                {'error': f'Неверное значение параметра days: {raw_days!r}'},
                status=400)
        days = max(0, min(days, 365))

    rows = expiring_batches(days=days)

    return Response({
        'отчёт': 'Партии с истекающим сроком годности',
        'дата': str(datetime.now().date()),
        'горизонт_дней': days,
        'просрочено': sum(1 for r in rows if r['expired']),
        'истекает': sum(1 for r in rows if not r['expired']),
        'строки': [{
            'товар': r['item_name'],
            'партия': r['batch_number'],
            'годен_до': str(r['expiry_date']),
            'дней_осталось': r['days_left'],
            'просрочено': r['expired'],
            'поступило': float(r['quantity_received']),
            'остаток_позиции': float(r['stock_remaining']),
            'склад': r['warehouse'],
            'поставщик': r['supplier'],
            'документ': r['document'],
        } for r in rows],
        'примечание': (
            'Остатки ведутся по позиции склада, а не по партиям, поэтому '
            'графа «остаток позиции» показывает общий остаток товара, а не '
            'то, сколько осталось именно от этой партии.'),
    })


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def production_cost_report(request):
    """Себестоимость выпущенной продукции: фактическая против плановой.

    Фактическая считается по последним закупочным ценам израсходованных
    материалов и включает только их — без зарплаты, амортизации и
    накладных расходов. Плановая берётся из карточки товара на момент
    выпуска.

    Параметр ?days= ограничивает период (по умолчанию 90 дней).
    """
    from warehouse.models import ProductionRun

    try:
        days = _int_param_simple(request.query_params.get('days'), 90)
    except ValueError as e:
        return Response({'error': str(e)}, status=400)

    since = timezone.localdate() - timedelta(days=days)
    runs = (ProductionRun.objects
            .filter(created_at__date__gte=since)
            .select_related('product', 'product_warehouse', 'created_by')
            .prefetch_related('materials__material'))

    rows = []
    for run in runs:
        deviation_percent = run.cost_deviation_percent
        rows.append({
            'номер': run.number,
            'дата': run.created_at.strftime('%d.%m.%Y'),
            'продукция': run.product.name,
            'артикул': run.product.article_number,
            'количество': float(run.quantity),
            'стоимость_материалов': float(run.material_cost),
            'себестоимость_факт': float(run.unit_cost),
            'себестоимость_план': float(run.planned_unit_cost),
            'отклонение': float(run.cost_deviation),
            'отклонение_процент': (round(float(deviation_percent), 1)
                                   if deviation_percent is not None else None),
            'цены_полные': run.pricing_complete,
            'материалы': [{
                'материал': line.material.name,
                'количество': float(line.quantity),
                'цена': float(line.unit_price),
                'сумма': float(line.total),
                'цена_известна': line.price_known,
            } for line in run.materials.all()],
        })

    over = [r for r in rows if r['отклонение'] > 0]
    incomplete = [r for r in rows if not r['цены_полные']]

    return Response({
        'отчёт': 'Себестоимость производства',
        'дата': str(timezone.localdate()),
        'период_дней': days,
        'выпусков': len(rows),
        'дороже_плана': len(over),
        'с_неполными_ценами': len(incomplete),
        'строки': rows,
        'примечание': (
            'Фактическая себестоимость включает только материалы по '
            'последним закупочным ценам. Зарплата, амортизация и накладные '
            'расходы в неё не входят.'),
    })


def _int_param_simple(value, default):
    if value in (None, ''):
        return default
    try:
        return max(1, min(int(value), 3650))
    except (TypeError, ValueError):
        raise ValueError(f'Неверное значение параметра days: {value!r}')
